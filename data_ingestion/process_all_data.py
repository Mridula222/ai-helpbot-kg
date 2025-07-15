import os
import requests
from bs4 import BeautifulSoup
import json
import pdfplumber
import docx
import openpyxl

# === 1. Scrape FAQs from MOSDAC ===
from selenium import webdriver
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.chrome.options import Options
from webdriver_manager.chrome import ChromeDriverManager
from bs4 import BeautifulSoup
import time

def scrape_mosdac_faq(url):
    options = Options()
    options.add_argument("--headless")
    options.add_argument("--disable-gpu")

    # ✅ Let webdriver-manager handle the driver path
    service = Service(ChromeDriverManager().install())
    driver = webdriver.Chrome(service=service, options=options)

    driver.get(url)
    time.sleep(3)

    soup = BeautifulSoup(driver.page_source, "html.parser")
    driver.quit()

    faqs = []
    question_tags = soup.find_all("h2", class_="panel-title")
    answer_blocks = soup.find_all("div", class_="field-items")

    for q, a in zip(question_tags, answer_blocks):
        faqs.append({
            "question": q.get_text(strip=True),
            "answer": a.get_text(strip=True)
        })

    return faqs

def save_faqs_to_txt_and_json(faqs):
    os.makedirs("data_ingestion/parsed_output", exist_ok=True)

    # Save as JSON
    with open("data_ingestion/mosdac_faqs.json", "w", encoding="utf-8") as f:
        json.dump(faqs, f, indent=4, ensure_ascii=False)

    # Save as .txt
    txt_content = ""
    for i, faq in enumerate(faqs, 1):
        txt_content += f"Q{i}: {faq['question']}\nA{i}: {faq['answer']}\n\n"

    with open("data_ingestion/parsed_output/mosdac_faqs.txt", "w", encoding="utf-8") as f:
        f.write(txt_content.strip())

# === 2. Parse PDF ===
def parse_pdf(file_path):
    text = ""
    with pdfplumber.open(file_path) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text += page_text + "\n\n"
    return text.strip()

# === 3. Parse DOCX ===
def parse_docx(file_path):
    doc = docx.Document(file_path)
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())

# === 4. Parse XLSX ===
def parse_xlsx(file_path):
    wb = openpyxl.load_workbook(file_path)
    text = ""
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            line = " | ".join(str(cell) if cell else "" for cell in row)
            text += line + "\n"
    return text.strip()

# === 5. Save individual .txt files ===
def save_text(text, filename):
    os.makedirs("data_ingestion/parsed_output", exist_ok=True)
    out_path = os.path.join("data_ingestion/parsed_output", filename)
    with open(out_path, "w", encoding="utf-8") as f:
        f.write(text)
    print(f"[+] Saved: {out_path}")

# === 6. Combine all .txt files into one master file ===
def combine_all_txt_to_master():
    output_dir = "data_ingestion/parsed_output"
    master_path = os.path.join(output_dir, "master_dataset.txt")

    with open(master_path, "w", encoding="utf-8") as master_file:
        for file in os.listdir(output_dir):
            if file.endswith(".txt") and file != "master_dataset.txt":
                file_path = os.path.join(output_dir, file)
                with open(file_path, "r", encoding="utf-8") as f:
                    content = f.read()
                    master_file.write(f"\n=== {file} ===\n")
                    master_file.write(content + "\n")
    
    print(f"\n Master file created at: {master_path}")

# === 7. Main pipeline ===
def process_all():
    print("Scraping FAQs from MOSDAC...")
    faqs = scrape_mosdac_faq("https://mosdac.gov.in/faq-page")

    if not faqs:
        print("[!] No FAQs found. Using dummy FAQs for now to continue the pipeline.")
        faqs = [
            {"question": "What is MOSDAC?", "answer": "MOSDAC stands for Meteorological and Oceanographic Satellite Data Archival Centre by ISRO."},
            {"question": "What satellites are supported?", "answer": "INSAT-3D, SCATSAT-1, Oceansat-2, Oceansat-3, etc."},
            {"question": "Is MOSDAC data free?", "answer": "Yes, most datasets are free to access after registration."},
            {"question": "What formats are used?", "answer": "NetCDF, HDF, GeoTIFF, and standard image formats."},
            {"question": "How can I register?", "answer": "Visit the registration section on MOSDAC and provide your credentials."}
        ]
    else:
        print(f"✅ Found {len(faqs)} FAQs")

    # ✅ Save FAQs (real or dummy)
    save_faqs_to_txt_and_json(faqs)


    print("\n Parsing local documents (PDF, DOCX, XLSX)...")
    docs_dir = "data_ingestion/docs"
    for file in os.listdir(docs_dir):
        path = os.path.join(docs_dir, file)
        name, ext = os.path.splitext(file)

        if ext.lower() == ".pdf":
            text = parse_pdf(path)
            save_text(text, f"{name}_pdf.txt")
        elif ext.lower() == ".docx":
            text = parse_docx(path)
            save_text(text, f"{name}_docx.txt")
        elif ext.lower() == ".xlsx":
            text = parse_xlsx(path)
            save_text(text, f"{name}_xlsx.txt")
        else:
            print(f"[!] Skipped unsupported file: {file}")

    print("\n Combining all .txt files into master_dataset.txt...")
    combine_all_txt_to_master()

if __name__ == "__main__":
    process_all()
